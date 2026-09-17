# -*- coding: utf-8 -*-
"""Generate HitRise software-copyright source-code DOCX.

The document follows the common PRC software-copyright source submission shape:
60 source pages, 50 non-empty source lines per page, using the first 30 pages
and last 30 pages of the selected Android APP source corpus.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(r"D:\2026\202605\hitrise")
ANDROID_ROOT = ROOT / "hitrise-android"
OUTPUT = ROOT / "hitrise-docs" / "HitRise_软件著作权源代码文档.docx"

LINES_PER_PAGE = 50
FRONT_PAGES = 30
BACK_PAGES = 30
TOTAL_SOURCE_PAGES = FRONT_PAGES + BACK_PAGES
SOURCE_LINE_TARGET = TOTAL_SOURCE_PAGES * LINES_PER_PAGE

SOURCE_ORDER = [
    "settings.gradle.kts",
    "build.gradle.kts",
    "app/build.gradle.kts",
    "app/proguard-rules.pro",
    "app/src/main/AndroidManifest.xml",
    "app/src/main/java/com/zclei/hitrise/MainActivity.kt",
    "app/src/main/java/com/zclei/hitrise/model/TrainingModels.kt",
    "app/src/main/java/com/zclei/hitrise/bluetooth/SensorBallBluetoothManager.kt",
    "app/src/main/java/com/zclei/hitrise/cloud/CloudModels.kt",
    "app/src/main/java/com/zclei/hitrise/cloud/CloudSyncService.kt",
    "app/src/main/java/com/zclei/hitrise/auth/ActivationModels.kt",
    "app/src/main/java/com/zclei/hitrise/auth/ActivationService.kt",
    "app/src/main/java/com/zclei/hitrise/AppPalette.kt",
    "app/src/main/java/com/zclei/hitrise/UiStrings.kt",
    "app/src/main/java/com/zclei/hitrise/ui/TrainingDashboardViews.kt",
    "app/src/main/java/com/zclei/hitrise/ui/CloudListAdapters.kt",
    "app/src/main/java/com/zclei/hitrise/ui/Haptics.kt",
    "app/src/main/java/com/zclei/hitrise/ui/Ripples.kt",
    "app/src/main/res/values/strings.xml",
    "app/src/main/res/values/themes.xml",
    "app/src/main/res/xml/file_paths.xml",
    "app/src/main/res/drawable/battery_status_background.xml",
    "app/src/main/res/drawable/ic_bluetooth_universal.xml",
    "app/src/main/res/drawable/ic_training_settings.xml",
]


def set_run_font(run, *, ascii_font: str, east_asia_font: str, size_pt: float) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run._element.rPr.rFonts.set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr_text, fld_sep, fld_text, fld_end])


def iter_source_files() -> Iterable[Path]:
    seen: set[Path] = set()
    for relative in SOURCE_ORDER:
        path = ANDROID_ROOT / relative
        if path.exists() and path.is_file():
            seen.add(path.resolve())
            yield path

    allowed_exts = {".kt", ".kts", ".xml", ".pro"}
    for path in sorted(ANDROID_ROOT.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in allowed_exts:
            continue
        parts = {part.lower() for part in path.parts}
        if {"build", ".gradle", "mipmap-hdpi", "mipmap-mdpi", "mipmap-xhdpi", "mipmap-xxhdpi", "mipmap-xxxhdpi"} & parts:
            continue
        if "drawable-nodpi" in parts:
            continue
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            yield path


def read_source_lines() -> list[str]:
    lines: list[str] = []
    forbidden_literals = ("Zz138575", "]v2Zu", "PRIVATE KEY", "BEGIN RSA")
    for path in iter_source_files():
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="utf-8-sig", errors="ignore")
        if any(secret in text for secret in forbidden_literals):
            continue
        for raw_line in text.splitlines():
            line = raw_line.rstrip().replace("\t", "    ")
            if line.strip():
                lines.append(line)
    if len(lines) < SOURCE_LINE_TARGET:
        return lines
    front_count = FRONT_PAGES * LINES_PER_PAGE
    back_count = BACK_PAGES * LINES_PER_PAGE
    return lines[:front_count] + lines[-back_count:]


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Consolas"
    normal.font.size = Pt(7)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_before = Pt(0)
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("HitRise Android APP Source Program Document")
    header_run.font.color.rgb = RGBColor(80, 96, 96)
    set_run_font(header_run, ascii_font="Calibri", east_asia_font="SimSun", size_pt=8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    prefix = footer.add_run("Page ")
    set_run_font(prefix, ascii_font="Calibri", east_asia_font="SimSun", size_pt=8)
    add_field(footer, "PAGE")
    suffix = footer.add_run(f" / {TOTAL_SOURCE_PAGES}")
    set_run_font(suffix, ascii_font="Calibri", east_asia_font="SimSun", size_pt=8)


def build_docx() -> tuple[Path, int]:
    lines = read_source_lines()
    if len(lines) < SOURCE_LINE_TARGET:
        raise RuntimeError(f"Not enough source lines: {len(lines)} < {SOURCE_LINE_TARGET}")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "HitRise Software Copyright Source Code"
    doc.core_properties.subject = "HitRise Android APP source program, 60 pages"
    doc.core_properties.author = "HitRise Project"
    doc.core_properties.comments = "Generated from Android source only; build outputs, media assets and runtime secrets excluded."

    selected = lines[:SOURCE_LINE_TARGET]
    for index, source_line in enumerate(selected):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(9)
        run = paragraph.add_run(source_line)
        set_run_font(run, ascii_font="Consolas", east_asia_font="SimSun", size_pt=7)
        if (index + 1) % LINES_PER_PAGE == 0 and index + 1 < len(selected):
            run.add_break(WD_BREAK.PAGE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT, len(selected)


if __name__ == "__main__":
    output, line_count = build_docx()
    print(f"Generated: {output}")
    print(f"Source lines: {line_count}")
    print(f"Source pages: {line_count // LINES_PER_PAGE}")
